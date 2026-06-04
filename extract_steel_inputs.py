"""
ecoinvent 鋼鐵製程技術投入萃取腳本 (多批次版本)
==============================================
功能：掃描多個 ecoinvent .spold 資料夾，找出所有鋼鐵相關製程，
      抓取其 technosphere inputs（技術投入），整理成 Excel。

使用方式：
  1. 從 EcoQuery > Files 下載整包資料庫（如 ecoinvent-3.12-cutoff-ecoSpold2.7z）
  2. 解壓縮後找到 datasets/ 資料夾（裡面有數萬個 .spold 檔案）
  3. 修改下方 DATASETS_FOLDERS 為你的所有路徑列表
  4. 執行：python extract_steel_inputs.py

需要安裝：
  pip install openpyxl lxml tqdm
"""

import os
import re
from pathlib import Path
from xml.etree import ElementTree as ET
from datetime import datetime
from collections import defaultdict

# ─── 設定區（請修改這裡）────────────────────────────────────────────
# 支援多個資料夾路徑（可一次處理多個批次）
DATASETS_FOLDERS = [
    r"C:\ecoinvent\batch_1\datasets",      # ← 批次 1
    r"C:\ecoinvent\batch_2\datasets",      # ← 批次 2
    # r"C:\ecoinvent\batch_3\datasets",    # ← 若有更多，取消註解並新增路徑
]

OUTPUT_FILE = r"steel_technosphere_inputs.xlsx"  # 輸出 Excel 檔名

# 搜尋關鍵字（小寫），符合其中一個就納入
STEEL_KEYWORDS = [
    "steel",
    "iron",
    "hot rolled",
    "cold rolled",
    "electric arc",
    "blast furnace",
    "pig iron",
    "cast iron",
    "steel production",
    "steel mill",
]

# 去重模式：
# - "merge"：相同製程只保留一筆（跨批次去重）
# - "keep_all"：所有批次的相同製程都保留，但標記批次來源
DEDUP_MODE = "merge"  # 改為 "keep_all" 若要保留所有批次的製程

# ──────────────────────────────────────────────────────────────────────

# ecoSpold2 XML namespace
NS = {
    "es2": "http://www.EcoInvent.org/EcoSpold02"
}


def is_steel_activity(activity_name: str, ref_product: str) -> bool:
    """判斷是否為鋼鐵相關製程"""
    combined = (activity_name + " " + ref_product).lower()
    return any(kw in combined for kw in STEEL_KEYWORDS)


def parse_spold(filepath: Path, batch_name: str = "") -> dict | None:
    """
    解析單一 .spold 檔案，回傳：
    {
        "activity_name": str,
        "ref_product": str,
        "geography": str,
        "batch": str,  # 新增：批次標識
        "file_path": str,  # 新增：原始檔案路徑
        "inputs": [ {"name": str, "unit": str, "amount": float, "comment": str} ]
    }
    若不是鋼鐵製程則回傳 None。
    """
    try:
        tree = ET.parse(filepath)
        root = tree.getroot()
    except ET.ParseError:
        print(f"  ⚠️  無法解析：{filepath.name}")
        return None

    # 取得 activity 基本資訊
    act_elem = root.find(".//es2:activity", NS)
    if act_elem is None:
        return None

    activity_name = act_elem.get("activityName", "")
    geography_elem = root.find(".//es2:geography/es2:shortname", NS)
    geography = geography_elem.text if geography_elem is not None else ""

    # 取得 reference product（output group=0）
    ref_product = ""
    for ie in root.findall(".//es2:intermediateExchange", NS):
        group = ie.find("es2:outputGroup", NS)
        if group is not None and group.text == "0":
            name_elem = ie.find("es2:name", NS)
            ref_product = name_elem.text if name_elem is not None else ""
            break

    if not is_steel_activity(activity_name, ref_product):
        return None

    # 抓 technosphere inputs（inputGroup = 5）
    inputs = []
    for ie in root.findall(".//es2:intermediateExchange", NS):
        group = ie.find("es2:inputGroup", NS)
        if group is None or group.text != "5":
            continue

        name_elem = ie.find("es2:name", NS)
        unit_elem  = ie.find("es2:unitName", NS)
        amount_str = ie.get("amount", "0")
        comment_elem = ie.find("es2:comment", NS)

        try:
            amount = float(amount_str)
        except ValueError:
            amount = 0.0

        inputs.append({
            "name":    name_elem.text  if name_elem    is not None else "",
            "unit":    unit_elem.text  if unit_elem    is not None else "",
            "amount":  amount,
            "comment": comment_elem.text if comment_elem is not None else "",
        })

    return {
        "activity_name": activity_name,
        "ref_product":   ref_product,
        "geography":     geography,
        "batch":         batch_name,  # 新增：記錄來自哪個批次
        "file_path":     str(filepath),
        "inputs":        inputs,
    }


def main():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from tqdm import tqdm

    # ──────────────────────────────────────────────────────────────────
    # 驗證所有資料夾
    # ──────────────────────────────────────────────────────────────────
    valid_folders = []
    for folder_path in DATASETS_FOLDERS:
        datasets_path = Path(folder_path)
        if datasets_path.exists():
            valid_folders.append(datasets_path)
            print(f"✅ 找到資料夾：{folder_path}")
        else:
            print(f"❌ 找不到資料夾：{folder_path}")

    if not valid_folders:
        print("\n❌ 沒有任何有效的資料夾！請檢查設定。")
        return

    # ──────────────────────────────────────────────────────────────────
    # 掃描所有 .spold 檔案
    # ──────────────────────────────────────────────────────────────────
    all_spold_files = []
    batch_info = {}  # 記錄每個批次的資料夾

    for batch_idx, datasets_path in enumerate(valid_folders, 1):
        batch_name = f"Batch_{batch_idx}"
        spold_files = list(datasets_path.glob("*.spold"))
        print(f"   🔍 {batch_name}：找到 {len(spold_files)} 個 .spold 檔案")
        all_spold_files.extend([(fp, batch_name) for fp in spold_files])
        batch_info[batch_name] = str(datasets_path)

    if not all_spold_files:
        print("❌ 沒有找到任何 .spold 檔案！")
        return

    print(f"\n✅ 總共找到 {len(all_spold_files)} 個檔案")
    print(f"🔍 開始篩選鋼鐵製程...\n")

    # ──────────────────────────────────────────────────────────────────
    # 解析所有檔案
    # ──────────────────────────────────────────────────────────────────
    all_results = []
    for fp, batch_name in tqdm(all_spold_files, desc="解析中"):
        result = parse_spold(fp, batch_name)
        if result is not None:
            all_results.append(result)

    print(f"✅ 找到 {len(all_results)} 個鋼鐵相關製程")

    if not all_results:
        print("❌ 沒有找到符合條件的製程，請確認關鍵字設定。")
        return

    # ──────────────────────────────────────────────────────────────────
    # 去重處理（若啟用 merge 模式）
    # ──────────────────────────────────────────────────────────────────
    if DEDUP_MODE == "merge":
        seen = set()
        deduplicated = []
        for res in all_results:
            key = (res["activity_name"], res["ref_product"], res["geography"])
            if key not in seen:
                seen.add(key)
                deduplicated.append(res)
        print(f"⚠️  去重後：{len(deduplicated)} 個獨特製程（移除 {len(all_results) - len(deduplicated)} 個重複）")
        all_results = deduplicated

    # ──────────────────────────────────────────────────────────────────
    # 建立 Excel
    # ──────────────────────────────────────────────────────────────────
    wb = openpyxl.Workbook()

    # ── Sheet 0：執行記錄 ──────────────────────────────────────────
    ws_log = wb.active
    ws_log.title = "執行記錄"

    ws_log["A1"] = "執行時間"
    ws_log["B1"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws_log["A2"] = "總製程數"
    ws_log["B2"] = len(all_results)
    ws_log["A3"] = "技術投入總筆數"
    ws_log["B3"] = sum(len(r["inputs"]) for r in all_results)
    ws_log["A4"] = "去重模式"
    ws_log["B4"] = DEDUP_MODE

    ws_log["A6"] = "批次來源"
    for batch_idx, (batch_name, folder) in enumerate(batch_info.items(), 1):
        ws_log[f"A{6 + batch_idx}"] = batch_name
        ws_log[f"B{6 + batch_idx}"] = folder

    for col in ["A", "B"]:
        ws_log.column_dimensions[col].width = 50

    # ── Sheet 1：製程總覽 ──────────────────────────────────────────
    ws_summary = wb.create_sheet("製程總覽")

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(color="FFFFFF", bold=True)
    headers = ["製程名稱", "參考產品", "地理位置", "來源批次", "技術投入數量"]

    for col, h in enumerate(headers, 1):
        cell = ws_summary.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, res in enumerate(all_results, 2):
        ws_summary.cell(row=row_idx, column=1, value=res["activity_name"])
        ws_summary.cell(row=row_idx, column=2, value=res["ref_product"])
        ws_summary.cell(row=row_idx, column=3, value=res["geography"])
        ws_summary.cell(row=row_idx, column=4, value=res["batch"])
        ws_summary.cell(row=row_idx, column=5, value=len(res["inputs"]))

    for col in ws_summary.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws_summary.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)

    # ── Sheet 2：技術投入明細 ──────────────────────────────────────
    ws_detail = wb.create_sheet("技術投入明細")

    detail_headers = ["製程名稱", "參考產品", "地理位置", "來源批次", "投入物質名稱", "單位", "數量", "備註"]
    for col, h in enumerate(detail_headers, 1):
        cell = ws_detail.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    row_idx = 2
    for res in all_results:
        for inp in res["inputs"]:
            ws_detail.cell(row=row_idx, column=1, value=res["activity_name"])
            ws_detail.cell(row=row_idx, column=2, value=res["ref_product"])
            ws_detail.cell(row=row_idx, column=3, value=res["geography"])
            ws_detail.cell(row=row_idx, column=4, value=res["batch"])
            ws_detail.cell(row=row_idx, column=5, value=inp["name"])
            ws_detail.cell(row=row_idx, column=6, value=inp["unit"])
            ws_detail.cell(row=row_idx, column=7, value=inp["amount"])
            ws_detail.cell(row=row_idx, column=8, value=inp["comment"])
            row_idx += 1

    for col in ws_detail.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws_detail.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)

    # ── Sheet 3：投入物質加總（跨所有製程）──────────────────────
    ws_agg = wb.create_sheet("投入物質加總")

    agg: dict[tuple, float] = defaultdict(float)
    for res in all_results:
        for inp in res["inputs"]:
            key = (inp["name"], inp["unit"])
            agg[key] += inp["amount"]

    agg_headers = ["投入物質名稱", "單位", "總量（所有鋼鐵製程加總）"]
    for col, h in enumerate(agg_headers, 1):
        cell = ws_agg.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for row_idx, ((name, unit), total) in enumerate(
        sorted(agg.items(), key=lambda x: -x[1]), 2
    ):
        ws_agg.cell(row=row_idx, column=1, value=name)
        ws_agg.cell(row=row_idx, column=2, value=unit)
        ws_agg.cell(row=row_idx, column=3, value=total)

    for col in ws_agg.columns:
        max_len = max(len(str(c.value or "")) for c in col)
        ws_agg.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)

    # 儲存
    wb.save(OUTPUT_FILE)
    print(f"\n🎉 完成！Excel 已儲存至：{OUTPUT_FILE}")
    print(f"   - 執行記錄：包含時間戳與批次信息")
    print(f"   - 製程總覽：{len(all_results)} 個製程")
    total_inputs = sum(len(r["inputs"]) for r in all_results)
    print(f"   - 技術投入明細：{total_inputs} 筆")
    print(f"   - 投入物質種類：{len(agg)} 種")


if __name__ == "__main__":
    main()
